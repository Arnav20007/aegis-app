from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).parent / "Aegis_Emergency_Response_Platform_Synopsis.docx"

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "EAF2F8"
GRAY = "D9E2F3"
TEXT = "1F2937"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def add_text(doc, text, size=11, bold=False, align=None, before=0, after=7, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=TEXT, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, size=15 if level == 1 else 12, bold=True, color=NAVY)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(item)
        set_font(run, size=10.5, color=TEXT)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_font(r, size=9.5, bold=True, color="FFFFFF")
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(value)
            set_font(r, size=9.2, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_page_number(section, page, total=11):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = "Aegis Emergency Response Platform Synopsis | Page {} of {}".format(page, total)
    for r in p.runs:
        set_font(r, size=8.5, color="6B7280")


def page_break(doc):
    doc.add_page_break()


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(11)
    return doc


def main():
    doc = setup_doc()

    # Page 1: cover page
    for _ in range(5):
        doc.add_paragraph()
    add_text(doc, "PROJECT SYNOPSIS", 15, True, WD_ALIGN_PARAGRAPH.CENTER, after=12)
    add_text(doc, "AEGIS EMERGENCY RESPONSE PLATFORM", 24, True, WD_ALIGN_PARAGRAPH.CENTER, after=10)
    add_text(doc, "A responsive web based safety monitoring and emergency response prototype", 13, False, WD_ALIGN_PARAGRAPH.CENTER, after=30, italic=True)
    add_text(doc, "Submitted in partial fulfillment of the requirements for the award of", 11, False, WD_ALIGN_PARAGRAPH.CENTER, after=3)
    add_text(doc, "Bachelor of Technology in Computer Science and Engineering", 13, True, WD_ALIGN_PARAGRAPH.CENTER, after=32)
    add_text(doc, "Submitted by", 11, False, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_text(doc, "[Student Name]", 15, True, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_text(doc, "Roll No. [Roll Number]", 11, False, WD_ALIGN_PARAGRAPH.CENTER, after=25)
    add_text(doc, "Under the guidance of", 11, False, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_text(doc, "[Guide Name]", 14, True, WD_ALIGN_PARAGRAPH.CENTER, after=32)
    add_text(doc, "Department of Computer Science and Engineering", 12, True, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_text(doc, "[College Name]", 13, True, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_text(doc, "[University Name]", 11, False, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_text(doc, "Academic Session 2026 - 2027", 11, False, WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_page_number(doc.sections[0], 1)
    page_break(doc)

    # Page 2: certificate and declaration
    add_heading(doc, "Certificate", 1)
    add_text(doc, "This is to certify that the project synopsis entitled Aegis Emergency Response Platform is a bona fide work proposed by [Student Name], Roll No. [Roll Number], of the Department of Computer Science and Engineering. The work is being carried out under the guidance of [Guide Name] as part of the minor project requirement for the Bachelor of Technology degree.")
    add_text(doc, "The proposed system is an academic prototype intended to demonstrate secure user access, emergency SOS activation, location-assisted monitoring, trusted-contact management, incident reporting, and responder communication. It is not presented as a substitute for official emergency services.")
    add_text(doc, "", after=26)
    add_table(doc, ["Signature", "Name and Designation"], [
        ["____________________________", "Project Guide\n[Guide Name]"],
        ["____________________________", "Head of Department\n[HOD Name]"],
        ["____________________________", "Student\n[Student Name]"],
    ], [3.2, 3.5])
    add_heading(doc, "Student Declaration", 1)
    add_text(doc, "I hereby declare that this synopsis is based on my proposed minor project work. The project has been prepared for academic evaluation and has not been submitted elsewhere for the award of any degree or diploma. All implementation claims in this document refer to the Aegis prototype developed for this project.")
    add_text(doc, "Place: [City]                                      Date: [Date]", 10.5, False, after=22)
    add_text(doc, "Signature of Student: ____________________________", 10.5, False, after=0)
    add_page_number(doc.sections[0], 2)
    page_break(doc)

    # Page 3: abstract
    add_heading(doc, "Abstract", 1)
    add_text(doc, "Personal safety applications must operate under stressful conditions where a person may have little time to navigate complex screens, communicate a precise location, or explain the situation. Aegis Emergency Response Platform is a responsive web based prototype designed to organize emergency response actions into one interface. The system provides a controlled SOS workflow, authenticated user access, trusted-contact management, live location monitoring, a fake call simulation, audio and image evidence capture, community hazard reporting, incident records, and a dispatcher style communication interface.")
    add_text(doc, "The project uses React with TypeScript for the client interface and FastAPI with SQLAlchemy for server-side services. The present development environment uses SQLite for convenient local testing; the architecture supports PostgreSQL through a configurable database URL for a production or multi-user deployment. JSON REST APIs are protected using JWT authentication for user-bound actions, while WebSocket communication supports real-time alert chat. OpenStreetMap is used for map visualization, and browser APIs are used for geolocation, media capture, vibration, clipboard, and sharing features where permission is granted.")
    add_text(doc, "The proposed system is structured as an emergency response platform rather than only a panic-button application. Its process can be summarized as emergency detection, alert generation, live monitoring, evidence collection, communication, and response tracking. The work demonstrates how a modular web application can combine user experience, security, data persistence, and browser capabilities for an academic safety solution.")
    add_heading(doc, "Keywords", 2)
    add_text(doc, "Women safety, emergency response, SOS alert, geolocation, FastAPI, React, JWT authentication, SQLAlchemy, SQLite, WebSocket, trusted contacts, OpenStreetMap.")
    add_heading(doc, "Project Profile", 2)
    add_table(doc, ["Item", "Description"], [
        ["Project Type", "Minor Project - Web Based Emergency Response Prototype"],
        ["Primary Users", "Individual users, trusted contacts, and simulated responders"],
        ["Main Goal", "Reduce the time required to activate and record an emergency response workflow"],
        ["Delivery Form", "Responsive browser application with REST APIs and WebSocket support"],
    ], [1.7, 5.0])
    add_page_number(doc.sections[0], 3)
    page_break(doc)

    # Page 4: introduction and problem
    add_heading(doc, "1 Introduction", 1)
    add_text(doc, "Safety concerns can arise while travelling, studying, working late, or moving through unfamiliar areas. During a stressful event, a user may be unable to type messages, locate a contact, remember a helpline, or convey accurate location details. Existing communication tools are useful, but their actions are usually distributed across different applications. A safety platform should make the most important actions easy to find, fast to trigger, and traceable after the event.")
    add_heading(doc, "1.1 Problem Statement", 2)
    add_text(doc, "The problem addressed by this project is the absence of a single, responsive academic prototype that combines emergency alert activation, authorized user data, location-assisted monitoring, trusted contacts, evidence tools, incident history, and responder communication in a coherent workflow. The system must be simple enough for rapid use but structured enough to preserve records and support a meaningful demonstration of emergency-response logic.")
    add_heading(doc, "1.2 Need for the System", 2)
    add_bullets(doc, [
        "A user needs an immediately visible SOS action with false-alarm mitigation.",
        "Trusted contacts need an organized location and alert context instead of fragmented communication.",
        "A responder view needs an alert status, last known coordinates, and a communication channel.",
        "Users may need supporting tools such as a fake call, media recording, camera capture, and location sharing.",
        "A project evaluator needs documented architecture, data flow, database entities, APIs, tests, limitations, and future scope.",
    ])
    add_heading(doc, "1.3 Proposed Solution", 2)
    add_text(doc, "Aegis provides a central dashboard that exposes safety functions through a mobile-friendly interface. The core workflow is deliberately staged: the user signs in, configures contacts, permits location access if desired, triggers SOS through a hold-and-confirm interaction, and monitors the resulting incident through archive and dispatcher views. The prototype stores records in a relational database and separates the frontend interface from the FastAPI backend so the modules can evolve independently.")
    add_page_number(doc.sections[0], 4)
    page_break(doc)

    # Page 5 objectives scope requirements
    add_heading(doc, "2 Objectives Scope and Requirements", 1)
    add_heading(doc, "2.1 Objectives", 2)
    add_bullets(doc, [
        "Design a responsive web interface for rapid emergency-oriented interactions.",
        "Implement registration, login, JWT-based access, and user profile retrieval.",
        "Create an SOS workflow that records an active alert and can later be resolved.",
        "Maintain trusted contacts, incident history, community hazards, and message records.",
        "Use browser geolocation, camera, microphone, sharing, and vibration features with user permission.",
        "Provide a responder-style console and WebSocket-based incident chat for demonstration.",
    ])
    add_heading(doc, "2.2 Scope", 2)
    add_text(doc, "The scope is a browser-based academic prototype. It demonstrates data capture, local persistence, role-oriented screens, and real-time communication concepts. The project does not directly integrate police dispatch systems, telecom SMS gateways, certified medical services, or a production-grade command center. Browser permissions, hardware support, connection quality, and operating-system settings influence certain features such as GPS, vibration, camera, microphone, and native share sheets.")
    add_heading(doc, "2.3 Functional Requirements", 2)
    add_table(doc, ["ID", "Requirement"], [
        ["FR-01", "Register, sign in, sign out, and retrieve an authenticated user profile."],
        ["FR-02", "Create, list, and remove trusted emergency contacts."],
        ["FR-03", "Trigger, list, update, and resolve emergency alerts."],
        ["FR-04", "Display live map telemetry and attempt browser geolocation with fallbacks."],
        ["FR-05", "Provide fake call, share location, audio recording, and camera capture tools."],
        ["FR-06", "Report and view community safety hazards."],
        ["FR-07", "Exchange alert-specific messages through a WebSocket channel."],
    ], [0.8, 5.9])
    add_heading(doc, "2.4 Non Functional Requirements", 2)
    add_text(doc, "The system should be responsive on laptop and mobile browsers, secure user-bound API calls with JWT, retain records reliably, provide understandable permission errors, protect sensitive data, and remain modular enough to support a PostgreSQL deployment, notification provider, and stronger authorization model in the future.")
    add_page_number(doc.sections[0], 5)
    page_break(doc)

    # Page 6 methodology architecture
    add_heading(doc, "3 Methodology and System Architecture", 1)
    add_text(doc, "The project follows an iterative implementation approach. Requirements were converted into small milestones: setup, authentication, profile data, trusted contacts, SOS alerting, location tracking, responder interface, notifications, background-monitoring concepts, and testing. Each module is independently testable through the UI and through REST API endpoints.")
    add_heading(doc, "3.1 Architectural Flow", 2)
    add_table(doc, ["Layer", "Responsibility", "Main Technologies"], [
        ["Presentation Layer", "Responsive screens, forms, dashboard, map view, media controls, and safety tools.", "React, TypeScript, Vite, browser APIs"],
        ["Application Layer", "Authentication, validation, alert processing, contact management, hazard records, and WebSocket handling.", "FastAPI, Python, JWT"],
        ["Persistence Layer", "Stores users, contacts, alerts, hazards, and chat messages.", "SQLAlchemy, SQLite for development, PostgreSQL-ready configuration"],
        ["Integration Layer", "Map display, browser geolocation, media capture, clipboard/share, and device vibration.", "OpenStreetMap, Geolocation API, MediaRecorder, getUserMedia"],
    ], [1.35, 3.35, 2.0])
    add_heading(doc, "3.2 Logical Data Flow", 2)
    add_text(doc, "User Browser -> React Interface -> FastAPI REST or WebSocket Endpoint -> SQLAlchemy ORM -> Relational Database. For maps and device functions, the React interface communicates with browser-controlled APIs after permission is granted. An SOS event creates or updates an alert record, refreshes the user archive, and exposes the incident to the dispatcher simulation.")
    add_heading(doc, "3.3 Emergency Response Flow", 2)
    add_text(doc, "User Action -> Hold SOS for 3 Seconds -> Five Second Cancellation Window -> Alert Created -> Coordinates Attached When Available -> Alert Archive and Dispatcher State Updated -> User and Operator Chat -> Incident Resolved -> Record Retained for Audit.")
    add_text(doc, "The staged SOS design is included to reduce accidental activation. The application separates confirmation from the final API call and provides a visual status change when an emergency becomes active.")
    add_page_number(doc.sections[0], 6)
    page_break(doc)

    # Page 7 modules
    add_heading(doc, "4 Major Modules", 1)
    add_table(doc, ["Module", "Purpose", "Key Behaviour"], [
        ["Authentication", "Identifies the user and protects account-specific functions.", "Registration, OAuth2-style form login, JWT storage, authenticated profile retrieval."],
        ["Trusted Contacts", "Maintains people who can be contacted during an emergency.", "Create, list, display priority, test fake-call interaction, and remove contact."],
        ["SOS and Alerts", "Creates the primary emergency record.", "Hold-to-arm interaction, cancellation interval, alert trigger, archive, and resolution state."],
        ["Live Tracking", "Presents current device location when permitted.", "Geolocation watch, relocation request, accuracy display, timeout fallback, and telemetry history."],
        ["Safety Tools", "Supplies supporting actions for an emergency situation.", "Fake call screen, native share or clipboard fallback, audio recording, camera capture."],
        ["Safety Hub", "Collects community context and emergency reference information.", "National helplines, hazard report form, and hazard registry."],
        ["Dispatcher and Chat", "Demonstrates responder-side monitoring and communication.", "Active alert list, logs, location context, WebSocket messages, and persisted chat history."],
        ["Audit and Analytics", "Provides a record for review and academic demonstration.", "Incident archive, medical profile, CSV export, summary counts, and visual indicators."],
    ], [1.35, 2.3, 3.05])
    add_heading(doc, "4.1 Browser Permission Strategy", 2)
    add_text(doc, "Location, microphone, camera, vibration, and sharing capabilities are permission-dependent. The frontend is designed to request them only when the corresponding feature is used. Failure states such as permission denial, unavailable hardware, and location timeout must be communicated to the user without crashing the application. High-accuracy GPS is allowed additional time, followed by a lower-accuracy network-location fallback where possible.")
    add_page_number(doc.sections[0], 7)
    page_break(doc)

    # Page 8 database schema
    add_heading(doc, "5 Database Design", 1)
    add_text(doc, "The backend uses SQLAlchemy ORM models. SQLite is used during development because it permits a self-contained demonstration database. The database connection is configurable through DATABASE_URL, making the model suitable for PostgreSQL in a later deployment. The schema is relational and uses foreign keys to associate contacts and alerts with a user, and chat messages with an alert.")
    add_heading(doc, "5.1 Entity Relationship Summary", 2)
    add_text(doc, "User 1:N Contact | User 1:N Alert | Alert 1:N ChatMessage | Hazard is a community record in the current prototype. A future version can associate each hazard with the reporting user and review status.")
    add_table(doc, ["Entity", "Important Fields", "Relationship / Purpose"], [
        ["User", "id, name, email, password_hash, blood_type, medical_notes, created_at", "Owns contacts and alerts; stores emergency medical profile."],
        ["Contact", "id, user_id, name, relation, phone, priority", "Trusted responder associated with one user."],
        ["Alert", "id, user_id, start_time, end_time, status, last_x, last_y", "Represents an SOS incident and last known coordinates."],
        ["ChatMessage", "id, alert_id, sender, text, timestamp", "Stores alert-specific messages exchanged through WebSocket."],
        ["Hazard", "id, category, description, x, y, reported_at", "Community safety observation displayed in the Safety Hub."],
    ], [1.15, 3.35, 2.2])
    add_heading(doc, "5.2 Data Protection Considerations", 2)
    add_text(doc, "Passwords are hashed using bcrypt through Passlib. JWT is used to identify users for protected actions. Future production work should move the JWT secret to environment variables, implement user and responder roles, authorize every alert operation, secure WebSockets, limit public data exposure, validate contact fields, and apply rate limits to community reporting endpoints.")
    add_page_number(doc.sections[0], 8)
    page_break(doc)

    # Page 9 API and tech
    add_heading(doc, "6 Technology Stack and API Design", 1)
    add_table(doc, ["Area", "Technology", "Reason for Selection"], [
        ["Frontend", "React with TypeScript", "Component-based UI and type checking for a responsive application."],
        ["Backend", "FastAPI with Python", "Fast API development, validation, async support, and automatic OpenAPI documentation."],
        ["ORM", "SQLAlchemy", "Maps Python models to relational database tables."],
        ["Authentication", "JWT and OAuth2 form login", "Supports protected endpoints without server-side session storage."],
        ["Maps", "OpenStreetMap", "Open map visualization without a mandatory commercial map API."],
        ["Real Time", "WebSocket", "Provides two-way alert-specific chat without polling."],
    ], [1.2, 2.1, 3.4])
    add_heading(doc, "6.1 REST API Summary", 2)
    add_table(doc, ["Method", "Endpoint", "Purpose"], [
        ["POST", "/auth/register", "Create a user profile."],
        ["POST", "/auth/login", "Authenticate and return a JWT access token."],
        ["GET", "/auth/me", "Retrieve the signed-in user profile."],
        ["GET / POST / DELETE", "/api/contacts/", "Manage trusted contacts."],
        ["POST", "/api/alerts/trigger", "Create or update the user's active SOS alert."],
        ["GET", "/api/alerts/my", "List the current user's incident archive."],
        ["POST", "/api/alerts/resolve/{id}", "Resolve an alert record."],
        ["GET / POST", "/api/hazards/", "View and report community hazards."],
        ["WebSocket", "/ws/alerts/{id}", "Exchange live alert chat messages."],
    ], [0.85, 2.5, 3.35])
    add_page_number(doc.sections[0], 9)
    page_break(doc)

    # Page 10 testing
    add_heading(doc, "7 Testing Results and Limitations", 1)
    add_heading(doc, "7.1 Verification Performed", 2)
    add_table(doc, ["Test Area", "Verification Result"], [
        ["Frontend build", "TypeScript validation and production Vite build completed without errors."],
        ["Backend health", "Health and OpenAPI endpoints returned successful responses after dependency setup."],
        ["Authentication", "Registration, login, JWT-protected profile retrieval, and sign-out flow were verified."],
        ["Records", "Trusted contact creation, SOS trigger, alert archive retrieval, alert resolution, and hazard reporting were verified."],
        ["UI navigation", "Dashboard, tracking, archive, contacts, safety hub, audit reports, analytics, and dispatcher screens rendered without frontend console errors."],
        ["Permission-sensitive features", "GPS, camera, microphone, ringtone, vibration, and share actions require final manual verification on the target laptop and mobile browser."],
    ], [1.7, 5.0])
    add_heading(doc, "7.2 Current Limitations", 2)
    add_bullets(doc, [
        "The prototype does not contact real emergency services, police, ambulance providers, or SMS gateways.",
        "Browser location and media functions depend on device support, browser permission, network quality, and user settings.",
        "SQLite is suited to local prototype testing; PostgreSQL is recommended for concurrent multi-user use.",
        "The current responder dashboard is a simulation and must not be interpreted as a live dispatch connection.",
        "Production deployment requires stricter access control, environment-managed secrets, WebSocket authorization, audit logging, and data-retention rules.",
    ])
    add_heading(doc, "7.3 Risk Mitigation", 2)
    add_text(doc, "The interface includes a hold-to-arm SOS design and a cancellation window to reduce accidental alerts. Location requests use longer high-accuracy timeouts and a network-location fallback. The application communicates failure states for unavailable browser features and keeps a persistent incident record for post-event review.")
    add_page_number(doc.sections[0], 10)
    page_break(doc)

    # Page 11 conclusion
    add_heading(doc, "8 Conclusion and Future Scope", 1)
    add_text(doc, "Aegis Emergency Response Platform demonstrates a complete academic safety workflow in a responsive web application. It combines secure account access, trusted contacts, SOS records, browser-assisted location monitoring, evidence tools, hazard reporting, audit information, and a responder-style interface. The design illustrates that a safety platform must focus on speed, clarity, consent, privacy, and traceability rather than only adding features.")
    add_heading(doc, "8.1 Future Scope", 2)
    add_bullets(doc, [
        "Integrate approved SMS, email, push-notification, and verified emergency contact channels.",
        "Add role-based access for users, trusted contacts, administrators, and verified responders.",
        "Migrate to PostgreSQL with automated migrations, backups, monitoring, and encrypted data-at-rest strategy.",
        "Add secure WebSocket authentication, message authorization, rate limiting, and incident audit logs.",
        "Introduce route safety scoring and risk analytics using verified, privacy-respecting incident data.",
        "Package the experience as a progressive web application with offline-safe screens and stronger mobile behavior.",
        "Develop consent-based AI assistance for risk signals while retaining human confirmation for emergency actions.",
    ])
    add_heading(doc, "References", 2)
    add_text(doc, "1. React Documentation. https://react.dev/\n2. FastAPI Documentation. https://fastapi.tiangolo.com/\n3. SQLAlchemy Documentation. https://docs.sqlalchemy.org/\n4. OpenStreetMap Documentation. https://www.openstreetmap.org/\n5. MDN Web Docs: Geolocation API, MediaDevices API, MediaRecorder API, Web Share API, and WebSocket API. https://developer.mozilla.org/", 9.8, False, after=10)
    add_heading(doc, "Appendix A Suggested Demonstration Sequence", 2)
    add_text(doc, "Sign in -> dashboard -> trusted contact -> live tracking permission -> fake call -> SOS hold and cancellation window -> active alert -> dispatcher view and chat -> resolve incident -> audit archive and CSV export.")
    add_page_number(doc.sections[0], 11)

    doc.core_properties.title = "Aegis Emergency Response Platform Synopsis"
    doc.core_properties.subject = "Minor Project Synopsis"
    doc.core_properties.author = "[Student Name]"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
